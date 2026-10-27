"""Document 06 acceptance tests (ZL-FSR-ENG-006 §18–19): semantic model,
HTML/Word/PDF/Excel rendering with read-back parity, audiences, issue,
filing refusal and external edits (publication.py, publication_store.py)."""
import io
from decimal import Decimal as D

import pytest
from sqlalchemy import select

from app.domains.audit_ledger.models import AuditEvent
from app.domains.financial_statements import assurance_store as qa, publication as pub, publication_store as ps, store
from app.domains.financial_statements.assembly import PackageInputs, Snapshot, assemble
from tests.test_financial_statements_assurance import (  # noqa: F401 - the db fixture is used by name
    LEAD, OFFICER, OUTSIDER, PREPARER, REVIEWER, _package, _ready, db)

CURRENT = {"bs.trade_receivables": "703000", "bs.cash": "512000", "bs.trade_payables": "-389000",
           "eq.share_capital": "-100000", "eq.retained_earnings": "-877000", "pl.revenue": "-3182000",
           "pl.cost_of_sales": "2431000", "pl.administrative_expenses": "902000"}
NOTE = {"N02": {"decision": "include", "text": "Revenue is recognised on delivery.", "evidence_ref": "p",
                "submitted_by": "a", "approved_by": "b"}}


def _model(audience="member_accounts", scale="units", state="WORKING_DRAFT", resolution=None, totals=CURRENT):
    snap = Snapshot("tb1", "sha256:a", True, {k: D(v) for k, v in totals.items()})
    package = {"id": "pkg1", "version": 1, "result": assemble(PackageInputs(snapshot=snap, disclosures=NOTE))}
    return pub.build_model(package=package, resolution=resolution, audience=audience, scale=scale, state=state)


# ── engine ──────────────────────────────────────────────────────────────────

def test_pub_t001_every_channel_reads_back_equal_to_the_model():
    model = _model()
    for channel, render in pub.RENDERERS.items():
        result = pub.parity(model, channel, render(model))
        assert result["ok"], (channel, result["issues"])
    assert model["model_hash"] == _model()["model_hash"]                                     # deterministic


def test_expenses_show_in_parentheses_and_missing_comparatives_are_not_zero():
    model = _model()
    pl = next(s for s in model["statements"] if s["statement_id"] == "profit_or_loss")
    cos = next(l for l in pl["lines"] if l["line_id"] == "cost_of_sales")
    assert cos["display"] == "(2,431,000)" and cos["amount"] == "2431000" and cos["display_sign"] == -1
    assert cos["prior_display"] == "not provided"                                              # PUB-T003
    assert pub.display(D("0"), "units") == "–" and pub.display(D("-1500"), "thousands") == "(2)"


def test_pub_t006_an_inverted_sign_in_the_pdf_fails_parity():
    model = _model()
    tampered = {**model, "statements": [{**s, "lines": [{**l, "display": l["display"].strip("()")} for l in s["lines"]]}
                                        for s in model["statements"]]}
    result = pub.parity(model, "pdf", pub.render_pdf(tampered))
    assert not result["ok"] and any("cost_of_sales" in i for i in result["issues"])


def test_pub_t011_a_screenshot_pdf_fails_semantic_and_accessibility_gates():
    from reportlab.pdfgen import canvas
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer)
    c.rect(50, 50, 400, 400, fill=1)                                                           # an image-like page, no text
    c.save()
    model = _model()
    assert not pub.parity(model, "pdf", buffer.getvalue())["ok"]
    assert "selectable text" in pub.accessibility("pdf", buffer.getvalue())["failed"]


def test_pub_t012_thousands_keep_exact_values_and_state_the_scale():
    model = _model(scale="thousands")
    revenue = next(l for s in model["statements"] for l in s["lines"] if l["concept_id"] == "profit_or_loss.revenue")
    assert revenue["display"] == "3,182" and revenue["amount"] == "3182000" and model["unit_label"] == "£000"
    assert pub.parity(model, "xlsx", pub.render_xlsx(model))["ok"]


def test_pub_t015_public_filleting_needs_a_small_company():
    big = _model(audience="public_filing", resolution={"result": {"size_class": "large"}, "facts": {}})
    assert big["blockers"] and any(s["statement_id"] == "profit_or_loss" for s in big["statements"])


def test_pub_t016_internal_ids_are_caught_in_public_copies():
    model = _model(audience="public_filing")
    leaked = pub.render_html({**model, "notes": [{"number": 1, "requirement_id": "N02", "title": "x",
                                                  "text": "Prepared by user-7f3a9 per workpaper W12, j.doe@firm.com"}]})
    result = pub.sensitive(model, "html", leaked, internal_ids=["user-7f3a9"])
    assert result["blocking"] and {"user-7f3a9", "j.doe@firm.com", "workpaper"} <= set(result["hits"])


def test_pub_t019_tagging_is_blocked_without_an_installed_licensed_taxonomy():
    result = pub.tagging(_model(), "UK_COMPANIES_HOUSE")
    assert not result["allowed"] and any("not installed" in r for r in result["reasons"])


def test_pub_t037_kpis_stay_out_of_the_statutory_statements():
    assert _model(audience="member_accounts")["analysis"] is None
    management = _model(audience="management_pack")
    assert "non-statutory" in management["analysis"]["label"]
    from openpyxl import load_workbook
    sheets = load_workbook(io.BytesIO(pub.render_xlsx(management))).sheetnames
    assert "Management analysis" in sheets and all(n.startswith(("Statutory", "Management", "Notes")) for n in sheets)


# ── stored workflow ─────────────────────────────────────────────────────────

async def _released(db):
    pkg, view = await _ready(db)
    view = await qa.attest(db, PREPARER, view["id"], role="preparer", scope="Prepared")
    view = await qa.attest(db, REVIEWER, view["id"], role="independent_reviewer", scope="Reviewed statements and notes")
    minute = next(e["id"] for e in view["evidence"] if e["filename"] == "board-minute.pdf")
    view = await qa.attest(db, OFFICER, view["id"], role="officer", scope="Approved", approval_evidence_id=minute)
    await qa.release(db, OFFICER, view["id"])
    return pkg


async def test_a_draft_renders_watermarked_and_cannot_be_issued(db):
    pkg = await _package(db)
    artifact = await ps.render(db, PREPARER, pkg["id"], audience="member_accounts", channel="pdf")
    assert artifact["state"] == "WORKING_DRAFT" and artifact["validation"]["parity"]["ok"]
    again = await ps.render(db, PREPARER, pkg["id"], audience="member_accounts", channel="pdf")
    assert again["id"] == artifact["id"]                                                        # PUB-T035 idempotent
    with pytest.raises(store.Refused, match="internal release"):                               # PUB-T002
        await ps.issue(db, OFFICER, pkg["id"], "member_accounts")


async def test_issue_after_release_statutory_stays_blocked_and_filing_is_never_sent(db):
    pkg = await _released(db)
    for channel in ("html", "docx", "pdf"):
        a = await ps.render(db, PREPARER, pkg["id"], audience="management_pack", channel=channel)
        assert a["state"] == "AUTHORISED_FOR_ISSUE" and a["validation"]["parity"]["ok"], a["validation"]
    with pytest.raises(store.Refused, match="did not prepare"):
        await ps.issue(db, PREPARER, pkg["id"], "management_pack")
    issued = await ps.issue(db, OFFICER, pkg["id"], "management_pack")
    assert issued["state"] == "ISSUED" and issued["release_id"].startswith("rel-") and len(issued["artifacts"]) == 3
    with pytest.raises(store.Refused, match="statutory claim"):
        await ps.issue(db, OFFICER, pkg["id"], "member_accounts")
    with pytest.raises(store.Refused, match="not authorised"):                                 # PUB-T038
        await ps.submit_filing(db, REVIEWER.__class__(id="x", tenant_id="t1", email="x@x", full_name="X", role="Learner"),
                               issued["id"], destination="UK_COMPANIES_HOUSE", live=True, confirmed=True)
    with pytest.raises(store.Refused, match="not certified"):                                  # PUB-T022/T028
        await ps.submit_filing(db, OFFICER, issued["id"], destination="UK_COMPANIES_HOUSE", live=True, confirmed=True)
    overview = await ps.overview(db, OFFICER, pkg["id"])
    assert overview["issues"][0]["filing"][0]["status"] == "NOT_SENT" and overview["issues"][0]["state"] == "ISSUED"
    names = [e.event_name for e in (await db.execute(select(AuditEvent))).scalars()]
    assert "financial_statements.publication_filing_unauthorised" in names


async def test_pub_t018_an_edited_word_file_is_quarantined_with_a_diff(db):
    pkg = await _package(db)
    artifact = await ps.render(db, PREPARER, pkg["id"], audience="member_accounts", channel="docx")
    from docx import Document
    _, _, content = await ps.download(db, PREPARER, artifact["id"])
    doc = Document(io.BytesIO(content))
    doc.tables[0].rows[1].cells[1].text = "9,999,999"
    buffer = io.BytesIO()
    doc.save(buffer)
    quarantined = await ps.reimport_docx(db, PREPARER, artifact["id"], buffer.getvalue())
    assert quarantined["state"] == "QUARANTINED" and quarantined["kind"] == "EXTERNAL_EDIT"
    assert any("9,999,999" in d for d in quarantined["validation"]["diff"])


async def test_pub_t031_another_tenant_cannot_download(db):
    pkg = await _package(db)
    artifact = await ps.render(db, PREPARER, pkg["id"], audience="member_accounts", channel="html")
    with pytest.raises(store.NotFound):
        await ps.download(db, OUTSIDER, artifact["id"])
    events = (await db.execute(select(AuditEvent).where(
        AuditEvent.event_name == "financial_statements.publication_access_denied"))).scalars().all()
    assert len(events) == 1 and events[0].tenant_id == "t2"
