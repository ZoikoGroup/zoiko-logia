"""Verified figures from attached documents' tables (orchestration/document_figures.py):
every file type, the real rejected financial-statements case, and the
service hook that always uses the complete documents."""
import io
from types import SimpleNamespace

import pytest

from app.domains.kriton_workspace.documents import extract_document
from app.orchestration import service
from app.orchestration.document_figures import tables_from_sources, verified_document_figures
from app.orchestration.figure_reconciliation import unreconciled_figures
from app.orchestration.websearch import WebSource

HEADER = "Customer\tCategory\tRevenue_USD\tCost_USD\tStatus"
ROWS = [
    "Summit Capital\tSoftware\t205,000\t67,000\tCompleted",
    "Pioneer Logistics\tHardware\t189,000\t90,000\tCompleted",
    "Horizon Bank\tSoftware\t176,000\t60,000\tIn Progress",
    "Nova Healthcare\tServices\t162,000\t80,000\tCompleted",
    "Kappa Manufacturing\tHardware\t156,000\t70,000\tCompleted",
]


def _src(snippet: str, title="sales.xlsx — 'Sales'!1:6") -> WebSource:
    return WebSource(title=title, url="", snippet=snippet)


def test_workbook_figures_cover_totals_groups_shares_top_n_and_remainders():
    block = verified_document_figures([_src("\n".join([HEADER, *ROWS]))])
    assert "Total Revenue_USD = 888,000" in block
    assert "Cost_USD as a share of Revenue_USD = 41.33%" in block
    assert "Revenue_USD by Category: Software = 381,000 (42.91%)" in block
    assert "Top 3 Customer by Revenue_USD together = 570,000 (64.19%); the other 2 together = 318,000 (35.81%)" in block
    assert "Rows by Status: Completed = 4; In Progress = 1" in block


def test_ranges_of_one_sheet_merge_and_the_repeated_header_and_total_rows_are_dropped():
    first = _src("\n".join([HEADER, *ROWS[:3]]), "sales.xlsx — 'Sales'!1:4")
    second = _src("\n".join([HEADER, *ROWS[3:], "Total\t\t888,000\t367,000\t"]), "sales.xlsx — 'Sales'!5:7")
    [table] = tables_from_sources([first, second])
    assert len(table.rows) == 6                         # 5 rows + the Total row, header dropped
    assert "Total Revenue_USD = 888,000" in verified_document_figures([first, second])   # Total row not double counted


def test_the_rejected_answer_now_reconciles_and_its_wrong_figure_does_not():
    # The real case: the model wrote "top 3 customers = 570,000 (25.2%)" and
    # "other 10 customers = 658,000". 570,000 / 25.2% are right but were
    # unverifiable; 658,000 was simply wrong (the other 10 total 722,000).
    rows = ["Summit Capital\t205000", "Pioneer Logistics\t189000", "Horizon Bank\t176000", "Nova Healthcare\t162000",
            "Kappa Manufacturing\t156000", "Core Systems\t145000", "Jupiter Energy\t134000", "River Utilities\t132000",
            "Apex Corp\t125000", "Fusion Labs\t118000", "Metro Telecom\t100000", "Delta Finance\t90000",
            "Alpha Retail\t85000", "Beta Media\t80000", "Gamma Foods\t75000", "Omega Travel\t70000",
            "Sigma Auto\t65000", "Theta Pharma\t60000", "Zeta Labs\t52000", "Lambda Group\t45000"]
    source = _src("\n".join(["Customer\tRevenue_USD", *rows]))
    evidence = source.snippet + verified_document_figures([source])
    assert unreconciled_figures("The top three customers contributed USD 570,000 (25.2 % of total revenue).", evidence) == []
    assert unreconciled_figures("The other 10 customers contributed 722,000.", evidence) == []
    assert unreconciled_figures("The other 10 customers contributed 658,000.", evidence) == ["658,000"]


def test_csv_word_and_powerpoint_tables_are_read_through_the_real_extractors():
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches

    csv_blocks = extract_document("ledger.csv", ("Account,Debit,Credit\nCash,1500,0\nSales,0,1200\nVAT,0,300\n").encode())
    doc = Document()
    doc.add_paragraph("Quarterly expenses")
    table = doc.add_table(rows=4, cols=2)
    for row, (label, amount) in zip(table.rows, [("Expense", "Amount"), ("Rent", "12,000"), ("Salaries", "48,000"), ("Utilities", "3,000")]):
        row.cells[0].text, row.cells[1].text = label, amount
    docx_bytes = io.BytesIO()
    doc.save(docx_bytes)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    grid = slide.shapes.add_table(4, 2, Inches(1), Inches(1), Inches(4), Inches(2)).table
    for r, (label, amount) in enumerate([("Region", "Sales"), ("North", "40,000"), ("South", "25,000"), ("West", "35,000")]):
        grid.cell(r, 0).text, grid.cell(r, 1).text = label, amount
    pptx_bytes = io.BytesIO()
    deck.save(pptx_bytes)

    for name, data, expected in (
        ("ledger.csv", None, "Total Debit = 1,500"),
        ("costs.docx", docx_bytes.getvalue(), "Total Amount = 63,000"),
        ("deck.pptx", pptx_bytes.getvalue(), "Total Sales = 100,000"),
    ):
        blocks = csv_blocks if data is None else extract_document(name, data)
        sources = [_src(b.text, f"{name} — {b.location}") for b in blocks]
        assert expected in verified_document_figures(sources), name


def test_pdf_text_tables_are_detected_from_label_and_number_lines():
    page = ("Income statement for FY2025\nRevenue by product\n"
            "Software 1,199,000 588,000\nHardware 553,000 280,000\nServices 512,000 241,000\n"
            "Prepared by finance.")
    block = verified_document_figures([_src(page, "report.pdf — Page 1")])
    assert "Total Column 1 = 2,264,000" in block and "Total Column 2 = 1,109,000" in block
    assert "Column 1 by Item: Software = 1,199,000 (52.96%)" in block


def test_prose_or_documents_without_numeric_tables_add_nothing():
    assert verified_document_figures([_src("This memo explains revenue recognition under IFRS 15.", "memo.docx — Document text")]) == ""
    assert verified_document_figures([]) == ""


def test_unrelated_columns_are_not_compared():
    block = verified_document_figures([_src("Customer\tRevenue\tUnits\nA\t100,000\t5\nB\t50,000\t3\nC\t80,000\t2")])
    assert "Units as a share of Revenue" not in block and "Total Units = 10" in block


def test_a_messy_sheet_with_titles_sections_formulas_dashes_codes_and_rates():
    sheet = "\n".join([
        "Quarterly sales report",
        "Prepared by\tFinance team",
        "Region\tRevenue\tCost\tProfit\tVAT rate",
        "North",                                              # a section row inside the table
        "Leeds\tUSD 1,000\t400\t=B4-C4 => 600\t20%",
        "York\t2,000\t-\t=B5-C5 => 2000\t20%",
        "South",
        "Bath\t3,000 USD\t(500)\t=B7-C7 => 3500\t5%",
        "Total\t6,000\t-100\t6,100\t",
    ])
    block = verified_document_figures([_src(sheet)])
    assert "Table 'sales.xlsx' (Region, Revenue, Cost, Profit, VAT rate): 3 data rows" in block   # title, "Prepared by" and Total rows left out
    assert "Total Revenue = 6,000" in block and "Total Cost = -100" in block and "Total Profit = 6,100" in block
    assert "VAT rate: average 15" in block and "Total VAT rate" not in block


def test_a_long_workbook_split_into_ranges_and_overlapping_chunks_is_counted_once():
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Customer", "Notes", "Revenue"])
    for n in range(1, 131):
        sheet.append([f"Customer {n:03d}", "long note " * 12, n * 100])
    data = io.BytesIO()
    workbook.save(data)
    blocks = extract_document("big.xlsx", data.getvalue())
    assert len({b.location for b in blocks}) >= 4 and len(blocks) > len({b.location for b in blocks})   # ranges, and split chunks
    sources = [WebSource(title=f"big.xlsx — {b.location}", url="", snippet=b.text, source_id="doc-1") for b in blocks]
    block = verified_document_figures(sources)
    assert f"Total Revenue = {sum(n * 100 for n in range(1, 131)):,}" in block
    assert "130 data rows" in block


def test_two_uploads_with_the_same_file_name_stay_separate():
    a = WebSource(title="sales.xlsx — 'S'!1:3", url="", snippet="Item\tAmount\nA\t10\nB\t20", source_id="doc-a")
    b = WebSource(title="sales.xlsx — 'S'!1:3", url="", snippet="Item\tAmount\nA\t1\nB\t2", source_id="doc-b")
    block = verified_document_figures([a, b])
    assert "Total Amount = 30" in block and "Total Amount = 3" in block


def test_pdf_column_names_come_from_a_header_line_or_a_year_row():
    named = verified_document_figures([_src("Product Revenue Cost\nSoftware 100 40\nHardware 50 20\nServices 30 10", "r.pdf — Page 1")])
    assert "Total Revenue = 180" in named and "Total Cost = 70" in named
    years = verified_document_figures([_src("Revenue 2024 2025\nSoftware 100 140\nHardware 50 60\nServices 30 35", "r.pdf — Page 2")])
    assert "Total 2025 = 235" in years


def test_ocr_boxes_become_tab_rows_with_empty_cells_kept_in_their_column():
    from app.domains.kriton_workspace.ocr import lines_from_boxes

    def box(x0, x1, y, text):
        return [[x0, y - 10], [x1, y - 10], [x1, y + 10], [x0, y + 10]], text, 0.95

    text = lines_from_boxes([
        box(50, 300, 20, "Sales summary FY2025"),
        box(50, 150, 80, "Product"), box(400, 520, 80, "Revenue"), box(700, 780, 80, "Cost"),
        box(50, 160, 120, "Software"), box(405, 530, 121, "1,199,000"), box(690, 800, 119, "588,000"),
        box(50, 160, 160, "Hardware"), box(690, 800, 160, "280,000"),            # Revenue cell blank
        box(50, 100, 200, "Data"), box(105, 160, 200, "Centre"), box(405, 530, 200, "512,000"), box(690, 800, 200, "241,000"),
    ])
    assert text.splitlines() == [
        "Sales summary FY2025", "Product\tRevenue\tCost", "Software\t1,199,000\t588,000",
        "Hardware\t\t280,000", "Data Centre\t512,000\t241,000",
    ]


def test_a_scanned_pdf_is_read_with_ocr():
    pytest.importorskip("rapidocr_onnxruntime")
    pytest.importorskip("pypdfium2")
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.load_default(size=30)
    image = Image.new("RGB", (1240, 700), "white")
    draw = ImageDraw.Draw(image)
    draw.text((80, 60), "Sales summary FY2025", fill="black", font=font)
    for r, row in enumerate([("Product", "Revenue", "Cost"), ("Software", "1,199,000", "588,000"),
                             ("Hardware", "553,000", "280,000"), ("Services", "512,000", "241,000")]):
        for c, cell in enumerate(row):
            draw.text((80 + c * 380, 160 + r * 70), cell, fill="black", font=font)
    scan = io.BytesIO()
    image.save(scan, "PDF", resolution=150)              # an image only: no text layer
    [page] = extract_document("scan.pdf", scan.getvalue())
    assert page.metadata == {"page": 1, "ocr": True}
    block = verified_document_figures([_src(page.text, "scan.pdf — Page 1")])
    assert "Total Revenue = 2,264,000" in block and "Revenue by Product: Software = 1,199,000 (52.96%)" in block


def test_ocr_unavailable_leaves_scanned_pages_empty_without_failing(monkeypatch):
    from PIL import Image

    from app.domains.kriton_workspace import ocr

    def unavailable():
        raise ImportError("libGL.so.1: cannot open shared object file")

    monkeypatch.setattr(ocr, "_get_engine", unavailable)
    scan = io.BytesIO()
    Image.new("RGB", (200, 200), "white").save(scan, "PDF")
    assert extract_document("scan.pdf", scan.getvalue()) == []


async def test_the_hook_reads_the_complete_documents_for_a_targeted_question(monkeypatch):
    full = [_src("\n".join([HEADER, *ROWS]))]
    partial = [_src("\n".join([HEADER, ROWS[0]]))]
    calls = []

    async def fake_retrieve(db, **kwargs):
        calls.append(kwargs)
        return full

    class _Nested:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *exc):
            return False

    monkeypatch.setattr(service, "retrieve_document_sources", fake_retrieve)
    db = SimpleNamespace(begin_nested=lambda: _Nested())
    targeted = SimpleNamespace(retrieval_mode="targeted")
    block = await service._verified_document_figures_context(db, "total revenue?", ["doc-1"], partial, targeted, "t1", "u1")
    assert "Total Revenue_USD = 888,000" in block and calls[0]["full_document"] is True
    assert await service._verified_document_figures_context(db, "q", [], partial, targeted, "t1", "u1") == ""


async def test_the_hook_fails_soft(monkeypatch):
    async def broken(db, **kwargs):
        raise TimeoutError("statement timeout")

    class _Nested:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *exc):
            return False

    monkeypatch.setattr(service, "retrieve_document_sources", broken)
    db = SimpleNamespace(begin_nested=lambda: _Nested())
    assert await service._verified_document_figures_context(
        db, "q", ["doc-1"], [], SimpleNamespace(retrieval_mode="targeted"), "t1", "u1") == ""


def test_a_fixed_asset_register_gets_its_straight_line_depreciation():
    register = "\n".join([
        "asset_id\tdescription\tcategory\tcost\tacquired\tuseful_life_years\taccumulated_depreciation\tnet_book_value",
        "FA201\tAsset 01\tPlant and machinery\t9450\t2023-02-15\t4\t4725\t4725",
        "FA202\tAsset 02\tMotor vehicles\t10900\t2023-03-15\t5\t6540\t4360",
        "FA203\tAsset 03\tIT equipment\t12350\t2023-04-15\t3\t4116.67\t8233.33",
        "FA204\tAsset 04\tPlant and machinery\t5520\t2023-05-15\t4\t2760\t2760",
    ])
    block = verified_document_figures([_src(register, "kriton-test-fixed-assets.xlsx — 'Register'!1:5")])
    # 9450/4 + 10900/5 + 12350/3 + 5520/4 = 2362.50 + 2180 + 4116.67 + 1380
    assert "Annual straight-line depreciation (cost ÷ useful_life_years" in block
    assert "total = 10,039.17 for 4 assets" in block
    assert "Plant and machinery = 3,742.50" in block and "IT equipment = 4,116.67" in block
    assert "FA201 = 2,362.50" in block
    assert "Total useful_life_years" not in block and "useful_life_years: average 4" in block


def test_a_residual_value_is_not_depreciated():
    register = "\n".join([
        "Asset ID\tDescription\tCost (GBP)\tUseful life (years)\tResidual value (GBP)",
        "FA-002\tForklift truck A\t38000\t8\t4000", "FA-008\tOffice laptops\t18000\t3\t0"])
    block = verified_document_figures([_src(register, "fa.xlsx — 'Fixed Assets'!1:3")])
    # (38,000 - 4,000) / 8 = 4,250 and 18,000 / 3 = 6,000
    assert "(Cost (GBP) − Residual value (GBP)) ÷ Useful life (years)" in block
    assert "FA-002 = 4,250" in block and "FA-008 = 6,000" in block and "total = 10,250 for 2 assets" in block


def test_no_depreciation_line_without_both_cost_and_useful_life():
    block = verified_document_figures([_src("Customer\tRevenue\tCost\nA\t100\t40\nB\t50\t20")])
    assert "straight-line" not in block


def test_a_dated_ledger_gets_totals_per_month_in_date_order():
    # "Show monthly totals as a bar chart" of a 600-row ledger was refused:
    # the figures had no month totals.
    rows = ["txn_id\tdate\tnominal_code\tamount",
            "T1\t2026-02-03\tNominal 4000\t100.50", "T2\t2026-01-15\tNominal 4010\t200",
            "T3\t2026-02-20\tNominal 4000\t50", "T4\t31/03/2026\tNominal 4010\t25"]
    figures = verified_document_figures([WebSource(title="ledger.csv — Sheet1", url="", snippet="\n".join(rows), source_id="d1")])
    assert "amount by month (from date): Jan 2026 = 200; Feb 2026 = 150.50; Mar 2026 = 25" in figures
    assert "Rows by month (from date): Jan 2026 = 1; Feb 2026 = 2; Mar 2026 = 1" in figures
    assert not unreconciled_figures("February 2026 total: 150.50", figures)


_GENERAL_LEDGER = ["Date\tVoucher No\tAccount\tDescription\tDebit\tCredit",
                   "2025-04-03\tJV-1\tTrade receivables\tSale\t1000\t", "2025-04-03\tJV-1\tRevenue\tSale\t\t1000",
                   "2025-04-28\tJV-2\tSalaries and wages\tPayroll\t300\t", "2025-04-28\tJV-2\tCash at bank\tPayroll\t\t300",
                   "2025-05-05\tJV-3\tCost of sales\tStock\t400\t", "2025-05-05\tJV-3\tTrade payables\tStock\t\t400",
                   "2025-05-10\tJV-4\tTrade receivables\tSale\t2500\t", "2025-05-10\tJV-4\tRevenue\tSale\t\t2500",
                   "2025-05-20\tJV-5\tCash at bank\tReceipt\t900\t", "2025-05-20\tJV-5\tTrade receivables\tReceipt\t\t900"]


def test_a_general_ledger_gets_income_and_expenses_per_month_not_trial_balance_statements():
    # "Show total revenue and expenses per month ... chart of monthly revenue"
    # was refused: Credit by month mixed every account, and the ledger's
    # lines were classified as a trial balance ("Cash at bank" on the credit
    # side stopped every total).
    figures = verified_document_figures(
        [WebSource(title="gl.csv — Sheet1", url="", snippet="\n".join(_GENERAL_LEDGER), source_id="d1")],
        "Show total revenue and expenses per month and draw a chart of monthly revenue")
    assert "Accounts not classified" not in figures and "Trial balance" not in figures
    assert "Revenue by month (Revenue): Apr 2025 = 1,000; May 2025 = 2,500" in figures
    assert "Total expenses by month: Apr 2025 = 300; May 2025 = 400" in figures
    assert "Profit (total income minus total expenses) by month: Apr 2025 = 700; May 2025 = 2,100" in figures
    assert "total income = 3,500; total expenses = 700; profit = 2,800" in figures
    assert "by month for Account" not in figures          # covered by the ledger lines; kept the request small
    assert not unreconciled_figures("Revenue was 2,500 in May 2025, against 1,000 in April 2025.", figures)


def test_a_dated_table_that_is_not_a_ledger_gets_months_per_customer():
    rows = ["Date\tCustomer\tAmount", "2026-01-05\tAlder\t100", "2026-01-09\tBirch\t40",
            "2026-02-03\tAlder\t70", "2026-02-11\tBirch\t10"]
    figures = verified_document_figures(
        [WebSource(title="sales.csv — Sheet1", url="", snippet="\n".join(rows), source_id="d1")], "Sales per month by customer")
    assert "Amount by month for Customer = Alder (from Date): Jan 2026 = 100; Feb 2026 = 70" in figures


def test_ledger_per_account_months_are_left_out_of_other_questions():
    figures = verified_document_figures(
        [WebSource(title="gl.csv — Sheet1", url="", snippet="\n".join(_GENERAL_LEDGER), source_id="d1")],
        "Which account has the largest debit total?")
    assert "by month for Account" not in figures and "Ledger 'gl.csv'" not in figures
    assert "Accounts not classified" not in figures      # still not read as a trial balance
    assert "Debit by month (from Date)" in figures        # the existing monthly totals stay


def _sheet(name: str, rows: list[str]) -> WebSource:
    return WebSource(title=f"pack.xlsx — '{name}'!1:{len(rows)}", url="", snippet="\n".join(rows), source_id="d1")


# Same column names on both sheets, with opposite meanings: a bank statement's
# Debit is money out, a cash book's Debit is money in.
_BANK = ["Date\tDetails\tDebit\tCredit\tBalance", "2026-09-01\tOpening balance\t\t\t1000",
         "2026-09-03\tReceipt - Alder Retail\t\t500\t1500", "2026-09-05\tCheque 101 - Fenwick\t200\t\t1300",
         "2026-09-12\tReceipt - Harbour Kitchens\t\t754\t2054", "2026-09-20\tDirect debit - Insurance\t100\t\t1954",
         "2026-09-30\tBank charges\t8\t\t1946", "2026-09-30\tInterest received\t\t4\t1950"]
_BOOK = ["Date\tDetails\tDebit\tCredit\tBalance", "2026-09-01\tOpening balance\t\t\t1000",
         "2026-09-03\tReceipt - Alder Retail\t500\t\t1500", "2026-09-05\tCheque 101 - Fenwick\t\t200\t1300",
         "2026-09-12\tReceipt - Harbour Kitchens\t745\t\t2045", "2026-09-29\tCheque 102 - Swift\t\t60\t1985",
         "2026-09-30\tReceipt - Elmfield (deposited)\t300\t\t2285"]


def test_a_bank_statement_and_cash_book_get_a_reconciliation():
    # A model-made reconciliation used the wrong closing balance, missed a
    # direct debit and called a cash-book error a bank error.
    figures = verified_document_figures([_sheet("Bank Statement", _BANK), _sheet("Cash Book", _BOOK)],
                                        "Prepare the bank reconciliation")
    assert "bank statement closing balance = 1,950; cash book closing balance = 2,285; difference = 335" in figures
    assert "receipts not yet credited (add to the bank balance): 2026-09-30, Receipt - Elmfield (deposited) 300" in figures
    assert "payments not yet presented (deduct from the bank balance): 2026-09-29, Cheque 102 - Swift 60" in figures
    assert "credits (add to the cash book): 2026-09-30, Interest received 4" in figures
    assert "debits (deduct from the cash book): 2026-09-20, Direct debit - Insurance 100; 2026-09-30, Bank charges 8" in figures
    assert "Receipt - Harbour Kitchens: cash book 745, bank 754, cash book understated by 9" in figures
    # 2,285 + 4 - 100 - 8 + 9 = 2,190 and 1,950 + 300 - 60 = 2,190
    assert "Adjusted cash book balance = 2,190; adjusted bank balance = 2,190; unexplained difference = 0" in figures
    assert not unreconciled_figures("The adjusted balance is 2,190 on both sides.", figures)


def test_a_running_balance_is_never_added_up():
    figures = verified_document_figures([_sheet("Bank Statement", _BANK)])
    assert "Total Balance" not in figures
    assert ("Balance is a running balance (never add it up): opening = 1,000 (2026-09-01, Opening balance); "
            "closing = 1,950 (2026-09-30, Interest received)") in figures
    assert "Bank reconciliation" not in figures          # a statement alone has nothing to reconcile


def test_an_undated_balance_column_is_still_totalled():
    figures = verified_document_figures([_src("Account\tBalance\nCash\t100\nBank\t250")])
    assert "Total Balance = 350" in figures


_VAT_ROWS = ["Date\tReference\tType\tCustomer/Supplier\tVAT treatment\tNet amount\tVAT amount",
             "2026-07-01\tS-1\tSale\tAlder\tStandard 20%\t1000\t200", "2026-07-02\tS-2\tSale\tBirch\tZero-rated\t500\t0",
             "2026-07-03\tS-3\tSale\tAlder\tStandard 20%\t300\t60", "2026-07-04\tP-1\tPurchase\tFenwick\tStandard 20%\t400\t80",
             "2026-07-05\tP-2\tPurchase\tOakline\tReduced 5%\t200\t10"]


def test_sales_by_vat_treatment_leave_out_the_purchases():
    # "Summarise sales by VAT treatment": the treatment totals mixed sales
    # and purchases, so the answer said the sources did not state it.
    figures = verified_document_figures([_src("\n".join(_VAT_ROWS), "vat.csv — Sheet1")],
                                        "Summarise sales by VAT treatment.")
    assert ("Net amount by VAT treatment for Type = Sale (total 1,800): Standard 20% = 1,300 (72.22%); "
            "Zero-rated = 500 (27.78%)") in figures
    assert "for Type = Purchase" not in figures          # not asked about
    assert "by Reference for Type" not in figures        # one row each says nothing
    assert "for Type = Sale" not in verified_document_figures([_src("\n".join(_VAT_ROWS), "vat.csv — Sheet1")],
                                                              "Total VAT by customer")


def test_a_trial_balance_gets_its_statements_and_the_suspense_account_is_flagged():
    # "Prepare the financial statements" put revenue and interest in the
    # balance sheet and stated no profit: the model classified the accounts.
    rows = ["account_code\taccount_name\tdebit\tcredit",
            "1000\tProperty, plant and equipment\t1265000\t", "1200\tInventory\t418000\t",
            "1300\tTrade receivables\t703000\t", "1400\tCash at bank\t512000\t",
            "2000\tTrade payables\t\t389000", "2200\tDeferred tax liability\t\t87000",
            "3000\tShare capital\t\t100000", "3100\tRetained earnings\t\t1000000",
            "4000\tRevenue\t\t4182000", "5000\tCost of sales\t2431000\t", "6000\tAdministrative expenses\t300000\t",
            "6100\tInterest payable\t62000\t", "8000\tSuspense\t67000\t"]
    figures = verified_document_figures([WebSource(title="tb.csv — Sheet1", url="", snippet="\n".join(rows), source_id="d1")])
    assert "total debits = 5,758,000; total credits = 5,758,000; difference = 0" in figures
    assert "Gross profit = 1,751,000" in figures and "Finance costs = 62,000 (Interest payable 62,000)" in figures
    assert "Profit for the year = 1,389,000" in figures
    assert "Total assets = 2,898,000" in figures and "Total liabilities = 476,000" in figures
    assert "Suspense account balance = 67,000" in figures
    assert "Total account_code" not in figures   # a code is not an amount


def test_contra_balances_and_common_accounts_still_give_the_statements():
    # Six ordinary accounts were "not classified", so no statement totals were
    # given and the income statement stopped after Revenue.
    rows = ["Account code\tAccount name\tDebit\tCredit",
            "1000\tLand and buildings\t400\t", "1090\tAccumulated depreciation\t\t100", "1200\tTrade receivables\t200\t",
            "1210\tAllowance for doubtful debts\t\t20", "3000\tShare capital\t\t300", "3200\tDividends paid\t50\t",
            "4000\tRevenue\t\t1000", "5000\tCost of sales\t600\t", "6010\tEmployer NI and pension\t40\t",
            "6600\tProfessional fees\t30\t", "2500\tBank loan (repayable 2030)\t\t100",
            # Balances the sheet (1,320 + 200 = 1,520): an unbalanced trial
            # balance gives no statements (Document 01 FS-009).
            "9999\tSuspense\t200\t", "\tTotal\t1520\t1520"]
    figures = verified_document_figures([WebSource(title="tb.csv — Sheet1", url="", snippet="\n".join(rows), source_id="d1")])
    assert "Accounts not classified" not in figures
    assert "Operating expenses = 70 (Employer NI and pension 40, Professional fees 30)" in figures
    assert "Profit for the year = 330" in figures
    assert "Non-current assets = 300 (Land and buildings 400, Accumulated depreciation -100)" in figures
    assert "Current assets = 180" in figures and "Non-current liabilities = 100" in figures
    assert "Dividends paid -50" in figures
    assert "Total = " not in figures                   # the sheet's own Total row is not an account


def test_an_account_the_rules_cannot_place_stops_the_statement_totals():
    rows = ["account_name\tdebit\tcredit", "Revenue\t\t500", "Mystery account\t500\t"]
    figures = verified_document_figures([WebSource(title="tb.csv — Sheet1", url="", snippet="\n".join(rows), source_id="d1")])
    assert "Accounts not classified" in figures and "Mystery account" in figures
    assert "Profit for the year" not in figures


def test_a_general_ledger_gets_each_account_balance_and_its_trial_balance():
    # "Generate ledger account statements": the model subtracted the wrong
    # debit and credit totals itself and the explanation was withheld.
    rows = ["Date\tVoucher No\tAccount\tDescription\tDebit\tCredit",
            "2025-04-03\tJV-1\tTrade receivables\tSale\t1000\t", "2025-04-03\tJV-1\tRevenue\tSale\t\t1000",
            "2025-04-20\tJV-2\tCash at bank\tReceipt\t600\t", "2025-04-20\tJV-2\tTrade receivables\tReceipt\t\t600",
            "2025-04-25\tJV-3\tRent\tRent\t250\t", "2025-04-25\tJV-3\tCash at bank\tRent\t\t250"]
    figures = verified_document_figures([WebSource(title="gl.csv — Sheet1", url="", snippet="\n".join(rows), source_id="d1")])
    assert "- Trade receivables: 2 entries; total debits = 1,000; total credits = 600; closing balance = 400 Dr" in figures
    assert "- Cash at bank: 2 entries; total debits = 600; total credits = 250; closing balance = 350 Dr" in figures
    assert "- Revenue: 1 entries; total debits = 0; total credits = 1,000; closing balance = 1,000 Cr" in figures
    assert "total debit balances = 1,000; total credit balances = 1,000; difference = 0" in figures
