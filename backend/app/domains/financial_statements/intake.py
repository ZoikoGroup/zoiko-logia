"""Receive and parse a trial-balance file (Document 04 stages P0–P1,
rules SD-001 to SD-003, SD-009, SD-012).

The original bytes are fingerprinted (SHA-256) and never modified. Active
content (macros) is rejected without being opened; external workbook links
are reported. Every sheet is read, hidden sheets and rows included, with each
cell's sheet, row and column kept so any amount can be traced to its cell.
Formulas are recorded with their cached values; nothing is recalculated.
"""
from __future__ import annotations

import csv
import hashlib
import io
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone

from .findings import HARD_STOP, INFO, REVIEW, Finding

MAX_BYTES = 20 * 1024 * 1024
_SPREADSHEET = (".xlsx",)
_MACRO_EXTENSIONS = (".xlsm", ".xlsb", ".xltm", ".xls")


@dataclass(frozen=True)
class SourceManifest:
    source_id: str
    sha256: str
    filename: str
    size_bytes: int
    received_at: str
    uploaded_by: str
    media_type: str
    sheets: tuple[str, ...] = ()
    hidden_sheets: tuple[str, ...] = ()

    def to_dict(self) -> dict:
        return {**self.__dict__, "sheets": list(self.sheets), "hidden_sheets": list(self.hidden_sheets)}


@dataclass(frozen=True)
class Cell:
    sheet: str
    row: int                       # 1-based, as the spreadsheet shows it
    column: int                    # 1-based
    text: str                      # the displayed / cached value as text
    formula: str | None = None
    hidden: bool = False           # on a hidden sheet or row

    @property
    def ref(self) -> str:
        return f"{self.sheet}!R{self.row}C{self.column}"


@dataclass
class ParsedSource:
    manifest: SourceManifest
    rows: list[list[Cell]] = field(default_factory=list)
    findings: list[Finding] = field(default_factory=list)
    rejected: bool = False


def _manifest(filename: str, content: bytes, uploaded_by: str, media_type: str,
              sheets=(), hidden=()) -> SourceManifest:
    digest = hashlib.sha256(content).hexdigest()
    return SourceManifest(
        source_id=f"src-{digest[:16]}", sha256=digest, filename=filename, size_bytes=len(content),
        received_at=datetime.now(timezone.utc).isoformat(), uploaded_by=uploaded_by, media_type=media_type,
        sheets=tuple(sheets), hidden_sheets=tuple(hidden),
    )


def _rejected(filename: str, content: bytes, uploaded_by: str, code: str, message: str, rule: str) -> ParsedSource:
    return ParsedSource(manifest=_manifest(filename, content, uploaded_by, "rejected"),
                        findings=[Finding(code, HARD_STOP, message, rule)], rejected=True)


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _parse_csv(filename: str, content: bytes, uploaded_by: str) -> ParsedSource:
    manifest = _manifest(filename, content, uploaded_by, "text/csv", sheets=("CSV",))
    parsed = ParsedSource(manifest=manifest)
    for row_number, values in enumerate(csv.reader(io.StringIO(_decode(content))), start=1):
        parsed.rows.append([Cell("CSV", row_number, column, (value or "").strip())
                            for column, value in enumerate(values, start=1)])
    return parsed


def _parse_xlsx(filename: str, content: bytes, uploaded_by: str) -> ParsedSource:
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = archive.namelist()
    except zipfile.BadZipFile:
        return _rejected(filename, content, uploaded_by, "UNREADABLE_FILE",
                         "The workbook is corrupted or not a valid .xlsx file.", "SD-012")
    if any(name.lower().endswith("vbaproject.bin") for name in names):
        return _rejected(filename, content, uploaded_by, "ACTIVE_CONTENT",
                         "The workbook contains macros; it was not opened. Save it as a plain .xlsx or .csv.", "SD-002")
    try:
        formulas = load_workbook(io.BytesIO(content), data_only=False, read_only=False)
        values = load_workbook(io.BytesIO(content), data_only=True, read_only=False)
    except Exception:  # noqa: BLE001 — any unreadable workbook is the same actionable error
        return _rejected(filename, content, uploaded_by, "UNREADABLE_FILE",
                         "The workbook could not be read.", "SD-012")
    hidden = [ws.title for ws in formulas.worksheets if ws.sheet_state != "visible"]
    manifest = _manifest(filename, content, uploaded_by,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                         sheets=[ws.title for ws in formulas.worksheets], hidden=hidden)
    parsed = ParsedSource(manifest=manifest)
    if any(name.startswith("xl/externalLinks/") for name in names):
        parsed.findings.append(Finding("EXTERNAL_LINKS", REVIEW,
                                       "The workbook links to other workbooks; linked values were not refreshed.", "SD-002"))
    for ws in formulas.worksheets:
        cached = values[ws.title]
        sheet_hidden = ws.sheet_state != "visible"
        hidden_rows = 0
        for row in ws.iter_rows():
            if not row:
                continue
            row_number = row[0].row
            row_hidden = sheet_hidden or bool(ws.row_dimensions[row_number].hidden)
            cells = []
            for cell in row:
                raw = cell.value
                formula = raw if isinstance(raw, str) and raw.startswith("=") else None
                shown = cached.cell(row=cell.row, column=cell.column).value if formula else raw
                cells.append(Cell(ws.title, cell.row, cell.column,
                                  "" if shown is None else str(shown).strip(), formula, row_hidden))
            if any(c.text for c in cells):
                parsed.rows.append(cells)
                hidden_rows += row_hidden and not sheet_hidden
        if sheet_hidden and any(c.text for r in parsed.rows for c in r if c.sheet == ws.title):
            parsed.findings.append(Finding("HIDDEN_CONTENT", REVIEW,
                                           f"Hidden sheet '{ws.title}' contains data; it is included and must be reviewed.",
                                           "SD-003", (ws.title,)))
        if hidden_rows:
            parsed.findings.append(Finding("HIDDEN_CONTENT", REVIEW,
                                           f"{hidden_rows} hidden row(s) on '{ws.title}' contain data; they are included.",
                                           "SD-003", (ws.title,)))
    if any(c.formula for r in parsed.rows for c in r):
        parsed.findings.append(Finding("FORMULAS_PRESENT", INFO,
                                       "Formula cells are read at their saved values; nothing was recalculated.", "SD-002"))
    return parsed


# Malware screening hook (Document 04 §15): a callable returning a reason to
# reject, or None. Unset, the macro, external-link and format controls above
# still apply; a deployment wires its scanning service here.
scan_upload = None


def receive(filename: str, content: bytes, uploaded_by: str) -> ParsedSource:
    """P0–P1: fingerprint, reject unsafe or unsupported files, read every cell."""
    name = (filename or "").lower()
    if scan_upload is not None and (reason := scan_upload(content)):
        return _rejected(filename, content, uploaded_by, "MALWARE_SUSPECTED", f"The file was rejected: {reason}.",
                         "SD-011")
    if len(content) > MAX_BYTES:
        return _rejected(filename, content, uploaded_by, "FILE_TOO_LARGE", "The file is larger than 20 MB.", "SD-012")
    if not content:
        return _rejected(filename, content, uploaded_by, "EMPTY_FILE", "The file is empty.", "SD-012")
    if name.endswith(_MACRO_EXTENSIONS):
        return _rejected(filename, content, uploaded_by, "ACTIVE_CONTENT",
                         "Macro-enabled or legacy workbooks are not opened. Save the trial balance as .xlsx or .csv.",
                         "SD-002")
    if name.endswith(".csv"):
        return _parse_csv(filename, content, uploaded_by)
    if name.endswith(_SPREADSHEET):
        return _parse_xlsx(filename, content, uploaded_by)
    return _rejected(filename, content, uploaded_by, "UNSUPPORTED_FORMAT",
                     "Upload the trial balance as .xlsx or .csv.", "SD-012")
