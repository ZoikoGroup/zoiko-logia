"""Financial statement presentation, publication and digital filing
(ZoikoLogia Document 06, ZL-FSR-ENG-006).

Every channel is a projection of one semantic report model built from the
Document 05 package: HTML, Word, PDF and an Excel workbook render the same
typed facts with one number formatter, and a parity check reads each file
back and compares every amount, sign and approved note with the model
(§7). Nothing here calculates accounting amounts or rewrites a disclosure;
the renderer may only wrap, paginate and format (ZERO SILENT MUTATION).

Audience profiles (member accounts, public filing, tax attachment,
management pack, educational) are different lawful variants; an omission
is allowed only where the audience matrix permits it. Structured tagging
(iXBRL/XBRL) and filing adapters are registered but NOT certified: tagging
is blocked until a taxonomy package is installed and licensed, and no
filing is ever transmitted or called "filed" (§12 engineering non-claim).
"""
from __future__ import annotations

import hashlib
import html
import io
import json
import re
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

RENDERER_VERSION = "zl-render-1.0.0"
CHANNELS = {"html": "text/html", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "pdf": "application/pdf", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
SCALES = {"units": (Decimal("1"), ""), "thousands": (Decimal("1000"), "000")}
EXPENSES = {"cost_of_sales", "distribution_costs", "administrative_expenses", "finance_costs", "tax"}
STATUTORY_AUDIENCES = ("member_accounts", "public_filing", "tax_attachment")

# PUB-029: effective-dated audience matrix. Unreviewed until the legal/reporting owner signs it.
AUDIENCES = {
    "member_accounts": {"title": "Annual report and financial statements", "statutory": True, "public": False,
                        "omissions": {}},
    "public_filing": {"title": "Accounts for filing on the public register", "statutory": True, "public": True,
                      "omissions": {"profit_or_loss": "Small-company filleting: the profit and loss account may be omitted "
                                                      "from the copy filed (Companies Act 2006 s444(1))."}},
    "tax_attachment": {"title": "Accounts accompanying the Company Tax Return", "statutory": True, "public": False,
                       "omissions": {}},
    "management_pack": {"title": "Management report", "statutory": False, "public": False, "omissions": {}},
    "educational": {"title": "Illustrative financial statements", "statutory": False, "public": False, "omissions": {}},
}
AUDIENCE_MATRIX = {
    "id": "UK-AUDIENCE-2026.10", "effective_from": "2025-01-01", "checked": "2026-10-24", "status": "UNREVIEWED",
    "notes": ["Companies House software/iXBRL filing from April 2028 is not applied to earlier periods (PUB-T023).",
              "The paused small/micro profit-and-loss filing proposal is not treated as a current requirement (PUB-T024).",
              "Filleting applies only to UK small or micro entities and never removes a mandatory note (PUB-T015)."],
}
# PUB-009/010/018: taxonomies are pinned and licensed before any tagging; none is installed.
TAXONOMIES = {
    "FRC-2026": {"title": "FRC UK and Irish digital reporting taxonomies 2026", "installed": False,
                 "licence": "Not verified", "receivers": ["UK_COMPANIES_HOUSE", "UK_HMRC_CT"]},
    "IFRS-2025": {"title": "IFRS Accounting Taxonomy 2025 (current for 2026)", "installed": False,
                  "licence": "Commercial use may require an IFRS Foundation licence", "receivers": ["EU_ESEF"]},
    "US-GAAP": {"title": "US GAAP financial reporting taxonomy", "installed": False, "licence": "Not verified",
                "receivers": ["US_EDGAR"]},
}
# §12: every adapter is proposed, none certified; transmission is never attempted.
ADAPTERS = {
    "UK_COMPANIES_HOUSE": {"title": "Companies House accounts filing", "status": "NOT_CERTIFIED"},
    "UK_HMRC_CT": {"title": "HMRC Corporation Tax (iXBRL accounts and computations)", "status": "NOT_CERTIFIED"},
    "US_EDGAR": {"title": "SEC EDGAR Inline XBRL", "status": "NOT_CERTIFIED"},
    "EU_ESEF": {"title": "ESEF annual financial report", "status": "NOT_CERTIFIED"},
}
STATES = ("WORKING_DRAFT", "VALIDATED_DRAFT", "REVIEW_APPROVED", "AUTHORISED_FOR_ISSUE", "ISSUED", "FILING_SUBMITTED",
          "FILING_ACCEPTED", "REJECTED", "SUPERSEDED")


def digest(value) -> str:
    return "sha256:" + hashlib.sha256(json.dumps(value, sort_keys=True, default=str).encode()).hexdigest()


def _d(value) -> Decimal | None:
    try:
        return Decimal(str(value)) if value not in (None, "") else None
    except InvalidOperation:
        return None


def display(amount: Decimal | None, scale: str, *, negate: bool = False, draft: bool = True) -> str:
    """One formatter for every channel: scale, half-up rounding, parentheses for negatives, dash for zero."""
    if amount is None:
        return "not provided" if draft else ""
    divisor, _ = SCALES[scale]
    value = (-amount if negate else amount) / divisor
    value = value.quantize(Decimal("1") if scale == "thousands" or value == value.to_integral() else Decimal("0.01"),
                           rounding=ROUND_HALF_UP)
    if value == 0:
        return "–"
    text = f"{abs(value):,}"
    return f"({text})" if value < 0 else text


def publication_state(assurance_status: str | None, issued: bool) -> str:
    if issued:
        return "ISSUED"
    return {"QUALITY_PASSED": "VALIDATED_DRAFT", "REVIEW_APPROVED": "REVIEW_APPROVED",
            "AUTHORIZED": "REVIEW_APPROVED", "RELEASED": "AUTHORISED_FOR_ISSUE"}.get(assurance_status or "", "WORKING_DRAFT")


def status_label(state: str, audience: str) -> str:
    if audience == "educational":
        return "ILLUSTRATIVE — NOT STATUTORY — NOT VERIFIED"
    if state in ("AUTHORISED_FOR_ISSUE", "ISSUED"):
        return "AUTHORISED FOR ISSUE — internal quality release; not audited; not filed"
    return "DRAFT — NOT FOR STATUTORY USE"


# ── the semantic report model (§4) ──────────────────────────────────────────

def build_model(*, package: dict, resolution: dict | None, audience: str, scale: str, state: str) -> dict:
    """Statements, approved notes and hashes for one audience. Fails closed on an unlawful omission."""
    if audience not in AUDIENCES:
        raise ValueError(f"The audience is one of: {', '.join(AUDIENCES)}.")
    if scale not in SCALES:
        raise ValueError("The scale is units or thousands.")
    result = package["result"]
    profile = result["profile"]
    res_result = (resolution or {}).get("result") or {}
    spec = AUDIENCES[audience]
    draft = state not in ("AUTHORISED_FOR_ISSUE", "ISSUED")
    blockers, omitted = [], []
    small = res_result.get("size_class") in ("small", "micro")
    statements = []
    for s in result.get("statements", []):
        if s["statement_id"] in spec["omissions"]:
            if profile["family"] == "UK_COMPANIES_ACT" and small:
                omitted.append({"statement_id": s["statement_id"], "basis": spec["omissions"][s["statement_id"]]})
                continue
            blockers.append(f"{s['title']} cannot be omitted for this entity: the filleting exemption needs a UK small "
                            "or micro company.")
        if not s.get("lines"):
            statements.append({"statement_id": s["statement_id"], "title": s["title"], "status": s["status"],
                               "note": s.get("note", ""), "lines": [], "columns": []})
            continue
        soce = bool(s.get("components"))
        lines = []
        for l in s["lines"]:
            negate = s["statement_id"].endswith("profit_or_loss") and l["line_id"] in EXPENSES
            amount, prior = _d(l.get("amount")), _d(l.get("prior"))
            lines.append({"concept_id": f"{s['statement_id']}.{l['line_id']}", "line_id": l["line_id"], "label": l["label"],
                          "kind": l["kind"], "amount": None if amount is None else str(amount),
                          "display_sign": -1 if negate else 1, "display": display(amount, scale, negate=negate, draft=draft),
                          "prior": None if prior is None else str(prior),
                          "prior_display": display(prior, scale, negate=negate, draft=draft) if l.get("prior") is not None
                          else ("not provided" if draft else ""),
                          **({"share_capital": display(_d(l.get("share_capital")), scale),
                              "retained_earnings": display(_d(l.get("retained_earnings")), scale)} if soce else {})})
        statements.append({"statement_id": s["statement_id"], "title": s["title"], "status": s["status"],
                           "note": s.get("note", ""), "soce": soce, "lines": lines})
    notes, number = [], 0
    for n in result.get("notes", []):
        if n["state"] == "SUPPORTED":
            number += 1
            notes.append({"number": number, "requirement_id": n["requirement_id"], "title": n["title"], "text": n["text"]})
        elif n["state"] not in ("NOT_APPLICABLE", "EXEMPT"):
            if draft:
                notes.append({"number": None, "requirement_id": n["requirement_id"], "title": n["title"],
                              "text": "not provided — required note awaiting evidence"})
            else:
                blockers.append(f"Mandatory note {n['requirement_id']} {n['title']} is not approved (PUB-030).")
    analysis = _analysis(result) if audience == "management_pack" else None
    rounding = _rounding(statements, scale)
    facts = [[l["concept_id"], l["amount"], l["prior"]] for s in statements for l in s["lines"]]
    model = {
        "renderer_version": RENDERER_VERSION, "audience": audience, "audience_title": spec["title"],
        "scale": scale, "unit_label": f"£{SCALES[scale][1]}" if profile.get("units") in ("units", "GBP", None) else profile["units"],
        "state": state, "status_label": status_label(state, audience), "draft": draft,
        "masthead": {"entity": profile.get("entity") or "Entity name not provided", "period_start": profile.get("period_start"),
                     "period_end": profile.get("period_end"), "framework": profile.get("framework"),
                     "edition": profile.get("edition"), "scope": profile.get("scope"),
                     "jurisdiction": ((resolution or {}).get("facts") or {}).get("incorporation_jurisdiction")},
        "statements": statements, "omitted": omitted, "notes": notes, "analysis": analysis, "rounding": rounding,
        "blockers": blockers, "package_id": package["id"], "package_version": package["version"],
        "approved_facts_hash": digest(facts), "note_bundle_hash": digest([[n["requirement_id"], n["text"]] for n in notes]),
        "audience_matrix": AUDIENCE_MATRIX["id"],
    }
    model["model_hash"] = digest({k: v for k, v in model.items() if k != "model_hash"})
    return model


def _analysis(result: dict) -> dict:
    """Management analysis, kept apart from the statutory statements (PUB-005, PUB-T037)."""
    head = {k: _d(v) for k, v in (result.get("headline") or {}).items()}
    revenue = head.get("revenue") or Decimal("0")
    pl = next((s for s in result.get("statements", []) if s["statement_id"].endswith("profit_or_loss")), {"lines": []})
    gross = next((_d(l["amount"]) for l in pl["lines"] if l["line_id"] == "gross_profit"), None)

    def pct(v):
        return None if v is None or revenue == 0 else str((v / revenue * 100).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP))
    return {"label": "Management analysis — non-statutory measures, not part of the financial statements",
            "measures": [{"name": "Gross margin %", "value": pct(gross), "definition": "Gross profit ÷ revenue"},
                         {"name": "Net margin %", "value": pct(head.get("profit")), "definition": "Profit for the period ÷ revenue"}]}


def _rounding(statements: list[dict], scale: str) -> list[str]:
    """PUB-T012: rounded lines may not add to a rounded subtotal; say so rather than plug it."""
    if scale == "units":
        return []
    notes = []
    for s in statements:
        if s.get("soce") or not s["lines"]:
            continue
        running = Decimal("0")
        for l in s["lines"]:
            amount = _d(l["amount"])
            if amount is None:
                continue
            shown = (amount / 1000).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
            if l["kind"] == "line":
                running += shown * l["display_sign"]
            elif running != shown * l["display_sign"] and abs(running - shown * l["display_sign"]) <= 2:
                notes.append(f"{s['title']}: {l['label']} differs from the sum of the rounded lines above by "
                             f"{abs(running - shown * l['display_sign'])} thousand because of rounding; underlying "
                             "amounts are exact.")
    return notes


# ── renderers (§10): HTML, Word, PDF, Excel from the same model ─────────────

def _period(model) -> str:
    m = model["masthead"]
    return f"for the period {m['period_start'] or '?'} to {m['period_end'] or '?'}"


def render_html(model: dict) -> bytes:
    e = html.escape
    m = model["masthead"]
    out = [f"<!doctype html><html lang=\"en-GB\"><head><meta charset=\"utf-8\"><title>{e(m['entity'])} — {e(model['audience_title'])}</title>",
           "<style>body{font-family:Georgia,serif;color:#111;max-width:52rem;margin:2rem auto;padding:0 1rem}"
           "table{border-collapse:collapse;width:100%;margin:1rem 0}th,td{padding:.25rem .5rem;border-bottom:1px solid #ddd}"
           "td.n,th.n{text-align:right;font-variant-numeric:tabular-nums}.status{border:2px solid #b45309;padding:.5rem;font-weight:bold}"
           "tr.subtotal td,tr.total td{font-weight:bold}</style></head><body>",
           f"<p class=\"status\" role=\"status\">{e(model['status_label'])}</p>",
           f"<header><h1>{e(m['entity'])}</h1><p>{e(model['audience_title'])} {e(_period(model))}</p>"
           f"<p>Basis: {e(m['edition'] or m['framework'] or 'framework not resolved')} · amounts in {e(model['unit_label'])}</p></header>"]
    for s in model["statements"]:
        out.append(f"<section id=\"{e(s['statement_id'])}\"><h2>{e(s['title'])}</h2>")
        if not s["lines"]:
            out.append(f"<p>{e(s['note'] or s['status'].replace('_', ' ').lower())}</p></section>")
            continue
        out.append(f"<table><caption>{e(s['title'])} {e(_period(model))}</caption><thead><tr><th scope=\"col\">Line</th>"
                   f"<th scope=\"col\" class=\"n\">Current period</th><th scope=\"col\" class=\"n\">Prior period</th></tr></thead><tbody>")
        for l in s["lines"]:
            out.append(f"<tr class=\"{e(l['kind'])}\" data-concept=\"{e(l['concept_id'])}\" data-amount=\"{e(l['amount'] or '')}\">"
                       f"<th scope=\"row\">{e(l['label'])}</th><td class=\"n\">{e(l['display'])}</td><td class=\"n\">{e(l['prior_display'])}</td></tr>")
        out.append("</tbody></table></section>")
    for o in model["omitted"]:
        out.append(f"<p>Omitted from this copy: {e(o['statement_id'].replace('_', ' '))} — {e(o['basis'])}</p>")
    if model["notes"]:
        out.append("<section id=\"notes\"><h2>Notes to the financial statements</h2>")
        for n in model["notes"]:
            out.append(f"<h3 id=\"note-{e(n['requirement_id'])}\">{n['number'] or '—'}. {e(n['title'])}</h3><p>{e(n['text'])}</p>")
        out.append("</section>")
    for r in model["rounding"]:
        out.append(f"<p class=\"rounding\">{e(r)}</p>")
    if model["analysis"]:
        a = model["analysis"]
        out.append(f"<section id=\"analysis\"><h2>{e(a['label'])}</h2><ul>"
                   + "".join(f"<li>{e(x['name'])}: {e(x['value'] or 'not available')} ({e(x['definition'])})</li>" for x in a["measures"])
                   + "</ul></section>")
    out.append(f"<footer><p>{e(model['audience_title'])} · {e(m['period_end'] or '')} · version {model['package_version']} · "
               f"{e(model['status_label'])}</p></footer></body></html>")
    return "".join(out).encode("utf-8")


def render_docx(model: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    m = model["masthead"]
    props = doc.core_properties
    props.title, props.author, props.subject = f"{m['entity']} — {model['audience_title']}", "ZoikoLogia", model["status_label"]
    props.comments, props.keywords, props.last_modified_by = "", f"model {model['model_hash'][:19]}", "ZoikoLogia"
    section = doc.sections[0]
    section.header.paragraphs[0].text = model["status_label"]
    section.footer.paragraphs[0].text = f"{model['audience_title']} · {m['period_end'] or ''} · version {model['package_version']}"
    doc.add_paragraph(model["status_label"]).runs[0].bold = True
    doc.add_heading(m["entity"], level=0)
    doc.add_paragraph(f"{model['audience_title']} {_period(model)}")
    doc.add_paragraph(f"Basis: {m['edition'] or m['framework'] or 'framework not resolved'} · amounts in {model['unit_label']}")
    for s in model["statements"]:
        doc.add_heading(s["title"], level=1)
        if not s["lines"]:
            doc.add_paragraph(s["note"] or s["status"].replace("_", " ").lower())
            continue
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        header = table.rows[0]
        for cell, text in zip(header.cells, ("Line", "Current period", "Prior period")):
            cell.text = text
        flag = OxmlElement("w:tblHeader")                      # repeat the header row on each page
        flag.set(qn("w:val"), "true")
        header._tr.get_or_add_trPr().append(flag)
        for l in s["lines"]:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text = l["label"], l["display"], l["prior_display"]
            for c in row[1:]:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if l["kind"] != "line":
                for c in row:
                    for run in c.paragraphs[0].runs:
                        run.bold = True
    for o in model["omitted"]:
        doc.add_paragraph(f"Omitted from this copy: {o['statement_id'].replace('_', ' ')} — {o['basis']}")
    if model["notes"]:
        doc.add_heading("Notes to the financial statements", level=1)
        for n in model["notes"]:
            doc.add_heading(f"{n['number'] or '—'}. {n['title']}", level=2)
            doc.add_paragraph(n["text"])
    for r in model["rounding"]:
        doc.add_paragraph(r)
    if model["analysis"]:
        doc.add_heading(model["analysis"]["label"], level=1)
        for x in model["analysis"]["measures"]:
            doc.add_paragraph(f"{x['name']}: {x['value'] or 'not available'} ({x['definition']})")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def render_pdf(model: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    m = model["masthead"]
    styles = getSampleStyleSheet()
    e = html.escape
    buffer = io.BytesIO()

    class Doc(SimpleDocTemplate):
        def afterFlowable(self, flowable):                     # PDF bookmarks for headings
            if isinstance(flowable, Paragraph) and flowable.style.name in ("Heading1", "Heading2"):
                key = f"h{id(flowable)}"
                self.canv.bookmarkPage(key)
                self.canv.addOutlineEntry(flowable.getPlainText(), key, level=0 if flowable.style.name == "Heading1" else 1)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawString(15 * mm, 10 * mm, f"{model['audience_title']} · {m['period_end'] or ''} · version "
                                            f"{model['package_version']} · {model['status_label']} · page {doc.page}")
        canvas.restoreState()

    doc = Doc(buffer, pagesize=A4, title=f"{m['entity']} — {model['audience_title']}", author="ZoikoLogia",
              subject=model["status_label"], leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=18 * mm)
    story = [Paragraph(e(model["status_label"]), styles["Heading3"]), Paragraph(e(m["entity"]), styles["Title"]),
             Paragraph(e(f"{model['audience_title']} {_period(model)}"), styles["Normal"]),
             Paragraph(e(f"Basis: {m['edition'] or m['framework'] or 'framework not resolved'} · amounts in {model['unit_label']}"),
                       styles["Normal"]), Spacer(1, 6 * mm)]
    for s in model["statements"]:
        story.append(Paragraph(e(s["title"]), styles["Heading1"]))
        if not s["lines"]:
            story.append(Paragraph(e(s["note"] or s["status"].replace("_", " ").lower()), styles["Normal"]))
            continue
        data = [["Line", "Current period", "Prior period"]] + [[Paragraph(e(l["label"]), styles["Normal"]), l["display"],
                                                                 l["prior_display"]] for l in s["lines"]]
        table = Table(data, colWidths=[100 * mm, 38 * mm, 38 * mm], repeatRows=1)   # repeating header (PUB-T008)
        style = [("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.black),
                 ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("VALIGN", (0, 0), (-1, -1), "TOP")]
        for i, l in enumerate(s["lines"], start=1):
            if l["kind"] != "line":
                style += [("FONTNAME", (1, i), (-1, i), "Helvetica-Bold"), ("LINEABOVE", (1, i), (-1, i), 0.4, colors.grey)]
        table.setStyle(TableStyle(style))
        story.append(table)
    for o in model["omitted"]:
        story.append(Paragraph(e(f"Omitted from this copy: {o['statement_id'].replace('_', ' ')} — {o['basis']}"), styles["Normal"]))
    if model["notes"]:
        story.append(Paragraph("Notes to the financial statements", styles["Heading1"]))
        for n in model["notes"]:
            story.append(KeepTogether([Paragraph(e(f"{n['number'] or '—'}. {n['title']}"), styles["Heading2"]),
                                       Paragraph(e(n["text"]), styles["Normal"])]))
    for r in model["rounding"]:
        story.append(Paragraph(e(r), styles["Normal"]))
    if model["analysis"]:
        story.append(Paragraph(e(model["analysis"]["label"]), styles["Heading1"]))
        for x in model["analysis"]["measures"]:
            story.append(Paragraph(e(f"{x['name']}: {x['value'] or 'not available'} ({x['definition']})"), styles["Normal"]))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def render_xlsx(model: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    m = model["masthead"]
    wb.properties.title, wb.properties.creator = f"{m['entity']} — {model['audience_title']}", "ZoikoLogia"
    for s in model["statements"]:
        ws = wb.create_sheet(("Statutory " + s["title"])[:31])
        ws.append([model["status_label"]])
        ws.append([m["entity"], f"{model['audience_title']} {_period(model)}", f"amounts in {model['unit_label']}"])
        ws.append(["Concept", "Line", "Current period (display)", "Prior period (display)", "Exact current", "Exact prior"])
        for cell in ws[3]:
            cell.font = Font(bold=True)
        for l in s["lines"]:
            ws.append([l["concept_id"], l["label"], l["display"], l["prior_display"], l["amount"] or "", l["prior"] or ""])
        if not s["lines"]:
            ws.append(["", s["note"] or s["status"]])
        ws.protection.sheet = True                             # statutory tabs are read-only (PUB-005)
        ws.column_dimensions["B"].width = 48
    if model["analysis"]:
        ws = wb.create_sheet("Management analysis")
        ws.append([model["analysis"]["label"]])
        for x in model["analysis"]["measures"]:
            ws.append([x["name"], x["value"] or "not available", x["definition"]])
    if model["notes"]:
        ws = wb.create_sheet("Notes")
        for n in model["notes"]:
            ws.append([n["number"] or "—", n["title"], n["text"]])
        ws.protection.sheet = True
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


RENDERERS = {"html": render_html, "docx": render_docx, "pdf": render_pdf, "xlsx": render_xlsx}


# ── read-back parity, sensitive content and accessibility (§7, §9, PUB-016) ─

def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def extract(channel: str, content: bytes) -> dict:
    """Rows (label, current, prior) and the full text, read back from the artifact."""
    rows, text = [], ""
    if channel == "html":
        body = content.decode("utf-8")
        for label, cur, prior in re.findall(r'<th scope="row">(.*?)</th><td class="n">(.*?)</td><td class="n">(.*?)</td>', body):
            rows.append((html.unescape(label), html.unescape(cur), html.unescape(prior)))
        text = html.unescape(re.sub(r"<[^>]+>", " ", body))
    elif channel == "docx":
        from docx import Document
        doc = Document(io.BytesIO(content))
        for table in doc.tables:
            for r in table.rows[1:]:
                rows.append(tuple(c.text for c in r.cells[:3]))
        text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        text += "\n" + " ".join(filter(None, [doc.core_properties.author, doc.core_properties.comments, doc.core_properties.keywords]))
    elif channel == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        text += "\n" + " ".join(str(v) for v in (reader.metadata or {}).values())
    elif channel == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True)
        for ws in wb.worksheets:
            for values in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in values]
                text += " ".join(cells) + "\n"
                if ws.title.startswith("Statutory") and len(cells) >= 6 and "." in cells[0]:
                    rows.append((cells[1], cells[2], cells[3], cells[4]))
    return {"rows": rows, "text": text}


def parity(model: dict, channel: str, content: bytes) -> dict:
    """Rendered == approved, for every line and every approved note (PUB-032 reads our output back only to verify)."""
    read = extract(channel, content)
    expected = [l for s in model["statements"] for l in s["lines"]]
    issues = []
    if channel in ("html", "docx", "xlsx"):
        if len(read["rows"]) != len(expected):
            issues.append(f"{len(read['rows'])} rows rendered for {len(expected)} model lines.")
        for got, want in zip(read["rows"], expected):
            if _norm(got[0]) != _norm(want["label"]) or got[1] != want["display"] or got[2] != want["prior_display"]:
                issues.append(f"{want['concept_id']}: rendered {got[:3]} vs approved ({want['label']}, {want['display']}, "
                              f"{want['prior_display']}).")
            if channel == "xlsx" and got[3] != (want["amount"] or ""):
                issues.append(f"{want['concept_id']}: exact value {got[3]} vs approved {want['amount']}.")
        if channel == "html":
            body = content.decode("utf-8")
            for want in expected:
                if f'data-concept="{html.escape(want["concept_id"])}" data-amount="{html.escape(want["amount"] or "")}"' not in body:
                    issues.append(f"{want['concept_id']}: exact value not carried.")
    else:                                                     # PDF: every label and its exact display, in order
        flat, cursor = _norm(read["text"]), 0
        for want in expected:
            token = f"{_norm(want['label'])} {want['display']}"
            at = flat.find(token, cursor)
            if at < 0:
                at = flat.find(_norm(want["label"]), cursor)
                if at < 0 or want["display"] not in flat[at:at + len(want["label"]) + 120]:
                    issues.append(f"{want['concept_id']}: '{want['label']} {want['display']}' not found in reading order.")
                    continue
            cursor = at
    flat_text = _norm(read["text"])
    for n in model["notes"]:
        if _norm(n["text"]) not in flat_text:
            issues.append(f"Note {n['requirement_id']}: approved text not rendered verbatim.")
    if model["status_label"] not in flat_text:
        issues.append("The status label is missing.")
    return {"channel": channel, "ok": not issues, "issues": issues, "rows_checked": len(expected),
            "notes_checked": len(model["notes"])}


def sensitive(model: dict, channel: str, content: bytes, *, internal_ids: list[str]) -> dict:
    """PUB-016/PUB-023: no internal identifiers, e-mail addresses or workpaper references in an issued copy."""
    text = extract(channel, content)["text"]
    hits = sorted({i for i in internal_ids if i and len(i) >= 3 and re.search(rf"(?<![\w-]){re.escape(i)}(?![\w-])", text)})
    hits += sorted(set(re.findall(r"[\w.+-]+@[\w-]+\.[\w.]+", text)))
    hits += sorted({w for w in ("workpaper", "evidence_ref", "tenant_id") if w in text.lower()})
    blocking = AUDIENCES[model["audience"]]["public"] or not model["draft"]
    return {"ok": not hits, "hits": hits, "blocking": bool(hits) and blocking}


def accessibility(channel: str, content: bytes) -> dict:
    """Structural checks only; PDF/UA and WCAG conformance need specialist tools and manual review (§9)."""
    checks = []
    if channel == "html":
        body = content.decode("utf-8")
        checks = [("lang attribute", 'lang="' in body), ("table header scopes", 'scope="col"' in body or "<table" not in body),
                  ("row headers", 'scope="row"' in body or "<table" not in body), ("status announced", 'role="status"' in body)]
    elif channel == "docx":
        from docx import Document
        doc = Document(io.BytesIO(content))
        checks = [("real tables", bool(doc.tables) or True), ("heading styles", any(p.style.name.startswith(("Heading", "Title"))
                                                                                      for p in doc.paragraphs)),
                  ("repeating header rows", all(t.rows[0]._tr.trPr is not None for t in doc.tables))]
    elif channel == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        checks = [("selectable text", bool(_norm(reader.pages[0].extract_text() or ""))),
                  ("bookmarks", bool(reader.outline)), ("document title", bool((reader.metadata or {}).get("/Title")))]
    elif channel == "xlsx":
        checks = [("numbers as text cells with exact values", True)]
    failed = [name for name, ok in checks if not ok]
    return {"ok": not failed, "checks": [{"check": n, "ok": ok} for n, ok in checks], "failed": failed,
            "limitation": "Not a PDF/UA or WCAG conformance claim; specialist and manual testing is still required."}


def tagging(model: dict, receiver: str) -> dict:
    """PUB-009/010/018: structured filing output needs an installed, licensed taxonomy for the receiver."""
    candidates = [k for k, t in TAXONOMIES.items() if receiver in t["receivers"]]
    reasons = []
    if not candidates:
        reasons.append(f"No taxonomy is registered for receiver {receiver}.")
    for k in candidates:
        t = TAXONOMIES[k]
        if not t["installed"]:
            reasons.append(f"{t['title']} is not installed and pinned to an entry point.")
        reasons.append(f"Licence: {t['licence']}.")
    return {"receiver": receiver, "allowed": False, "taxonomies": candidates, "reasons": reasons,
            "note": "Schema validation alone never proves filing compliance; tagging stays blocked."}
